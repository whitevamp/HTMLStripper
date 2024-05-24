import os
import tkinter as tk
from tkinter import filedialog, ttk
from bs4 import BeautifulSoup
import docx
import pandas as pd
from fpdf import FPDF
from ebooklib import epub

# Function to strip HTML and JavaScript
def strip_html_js(content):
    # print("Stripping HTML and JavaScript...")
    soup = BeautifulSoup(content, 'html.parser')
    for script in soup(["script", "style"]):
        script.decompose()
    return soup.get_text()

# Function to save content in different formats
def save_content(content, output_path, output_format):
    # print(f"Saving content to {output_path} as {output_format}...")
    if output_format == 'txt':
        with open(output_path, 'w', encoding='utf-8') as file:
            file.write(content)
    elif output_format == 'doc':
        doc = docx.Document()
        doc.add_paragraph(content)
        doc.save(output_path)
    elif output_format == 'csv':
        df = pd.DataFrame([line.split() for line in content.split('\n')])
        df.to_csv(output_path, index=False)
    elif output_format == 'pdf':
        pdf = FPDF()
        pdf.add_page()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.set_font("Arial", size=12)
        for line in content.split('\n'):
            try:
                pdf.cell(200, 10, txt=line.encode('latin-1', 'replace').decode('latin-1'), ln=True)
            except UnicodeEncodeError:
                pass
        pdf.output(output_path)
    elif output_format == 'epub':
        book = epub.EpubBook()
        book.set_identifier('id123456')
        book.set_title('Sample')
        book.set_language('en')
        chapter = epub.EpubHtml(title='Content', file_name='chap_01.xhtml', lang='en')
        chapter.content = f'<html><body>{content}</body></html>'
        book.add_item(chapter)
        book.spine = ['nav', chapter]
        epub.write_epub(output_path, book)
    elif output_format == 'html':
        with open(output_path, 'w', encoding='utf-8') as file:
            file.write(content)

# Recursive function to process HTML files
def process_directory(input_directory, output_directory, output_format):
    # print(f"Processing directory: {input_directory}")
    for root, _, files in os.walk(input_directory):
        # print(f"Checking directory: {root}")
        for file in files:
            # print(f"Found file: {file}")
            if file.endswith(('.html', '.htm')):
                input_path = os.path.join(root, file)
                output_path = os.path.join(output_directory, os.path.splitext(file)[0] + '.' + output_format)
                # print(f"Processing file: {input_path}")
                with open(input_path, 'r', encoding='utf-8') as f:
                    content = f.read()
                stripped_content = strip_html_js(content)
                save_content(stripped_content, output_path, output_format)
    # print("Processing complete.")

# Tkinter UI
def select_input_directory():
    directory = filedialog.askdirectory()
    if directory:
        input_directory_var.set(directory)
        # print(f"Selected input directory: {directory}")

def select_output_directory():
    directory = filedialog.askdirectory()
    if directory:
        output_directory_var.set(directory)
        # print(f"Selected output directory: {directory}")

def start_processing():
    input_directory = input_directory_var.get()
    output_directory = output_directory_var.get()
    output_format = format_var.get()
    if input_directory and output_directory and output_format:
        # print(f"Starting processing with input directory: {input_directory}, output directory: {output_directory}, and format: {output_format}")
        process_directory(input_directory, output_directory, output_format)
    else:
        # print("Please select both input and output directories, and the output format.")
        pass

root = tk.Tk()
root.title("HTML Stripper")

input_directory_var = tk.StringVar()
output_directory_var = tk.StringVar()
format_var = tk.StringVar()

tk.Label(root, text="Select Input Directory:").grid(row=0, column=0, padx=10, pady=10)
tk.Entry(root, textvariable=input_directory_var, width=50).grid(row=0, column=1, padx=10, pady=10)
tk.Button(root, text="Browse", command=select_input_directory).grid(row=0, column=2, padx=10, pady=10)

tk.Label(root, text="Select Output Directory:").grid(row=1, column=0, padx=10, pady=10)
tk.Entry(root, textvariable=output_directory_var, width=50).grid(row=1, column=1, padx=10, pady=10)
tk.Button(root, text="Browse", command=select_output_directory).grid(row=1, column=2, padx=10, pady=10)

tk.Label(root, text="Output Format:").grid(row=2, column=0, padx=10, pady=10)
format_options = ['txt', 'doc', 'csv', 'pdf', 'epub', 'html']
ttk.Combobox(root, textvariable=format_var, values=format_options).grid(row=2, column=1, padx=10, pady=10)

tk.Button(root, text="Start", command=start_processing).grid(row=3, column=0, columnspan=3, pady=20)

root.mainloop()
